"""Resume parser supporting PDF, DOCX, and TXT files with robust error handling."""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Optional, Union
import docx
import pdfplumber


@dataclass
class ResumeParseResult:
    """Structured result returned by the resume parser."""
    text: str
    success: bool
    error_message: Optional[str]
    file_type: str
    word_count: int
    char_count: int


def _extract_from_pdf(file_source: Union[str, Path, BytesIO]) -> str:
    """Extract text from a PDF document using pdfplumber."""
    text_chunks = []
    with pdfplumber.open(file_source) as pdf:
        if not pdf.pages:
            raise ValueError("The PDF document contains no pages.")
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_chunks.append(page_text)

    extracted = "\n".join(text_chunks).strip()
    if not extracted:
        raise ValueError(
            "No extractable text found. The PDF may be scanned, image-only, or password-protected. "
            "OCR is not enabled; please supply a text-based PDF, DOCX, or TXT file."
        )
    return extracted


def _extract_from_docx(file_source: Union[str, Path, BytesIO]) -> str:
    """Extract text from a DOCX document using python-docx."""
    doc = docx.Document(file_source)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text and cell_text not in paragraphs:
                    paragraphs.append(cell_text)

    extracted = "\n".join(paragraphs).strip()
    if not extracted:
        raise ValueError("The DOCX document is empty or contains no readable text paragraphs.")
    return extracted


def _extract_from_txt(file_source: Union[str, Path, BytesIO]) -> str:
    """Extract text from a plain text file using common encodings."""
    if isinstance(file_source, BytesIO):
        raw_bytes = file_source.getvalue()
    elif hasattr(file_source, "read"):
        raw_bytes = file_source.read()
    else:
        path = Path(file_source)
        raw_bytes = path.read_bytes()

    if not raw_bytes:
        raise ValueError("The TXT file is completely empty (0 bytes).")

    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252", "utf-16"):
        try:
            decoded = raw_bytes.decode(encoding).strip()
            if decoded:
                return decoded
        except UnicodeDecodeError:
            continue

    raise ValueError("Unable to decode the text file using standard character encodings.")


def parse_resume(
    file_source: Union[str, Path, BytesIO],
    filename: Optional[str] = None
) -> ResumeParseResult:
    """Parse resume from a file path or in-memory BytesIO buffer.
    
    Parameters
    ----------
    file_source : Union[str, Path, BytesIO]
        Path to file or file-like object (e.g. from st.file_uploader).
    filename : Optional[str]
        Explicit filename (useful when passing BytesIO), to detect extension.
        
    Returns
    -------
    ResumeParseResult
        Object containing extracted text, success flag, word count, and error if any.
    """
    if file_source is None:
        return ResumeParseResult(
            text="",
            success=False,
            error_message="No file source provided.",
            file_type="unknown",
            word_count=0,
            char_count=0,
        )

    # Determine extension
    ext = ""
    if filename:
        ext = Path(filename).suffix.lower()
    elif isinstance(file_source, (str, Path)):
        ext = Path(file_source).suffix.lower()

    # If it's a file path, verify it exists and is not empty
    if isinstance(file_source, (str, Path)):
        path = Path(file_source)
        if not path.exists():
            return ResumeParseResult(
                text="",
                success=False,
                error_message=f"File does not exist at path: {path}",
                file_type=ext.replace(".", ""),
                word_count=0,
                char_count=0,
            )
        if path.stat().st_size == 0:
            return ResumeParseResult(
                text="",
                success=False,
                error_message="The uploaded file is empty (0 bytes).",
                file_type=ext.replace(".", ""),
                word_count=0,
                char_count=0,
            )

    try:
        if ext == ".pdf":
            text = _extract_from_pdf(file_source)
            file_type = "pdf"
        elif ext in (".docx", ".doc"):
            if ext == ".doc":
                return ResumeParseResult(
                    text="",
                    success=False,
                    error_message="Legacy .doc format is not supported. Please convert to modern .docx or .pdf.",
                    file_type="doc",
                    word_count=0,
                    char_count=0,
                )
            text = _extract_from_docx(file_source)
            file_type = "docx"
        elif ext in (".txt", ".text"):
            text = _extract_from_txt(file_source)
            file_type = "txt"
        else:
            return ResumeParseResult(
                text="",
                success=False,
                error_message=f"Unsupported file format '{ext}'. Only PDF, DOCX, and TXT are supported.",
                file_type=ext.replace(".", "") or "unknown",
                word_count=0,
                char_count=0,
            )

        words = text.split()
        return ResumeParseResult(
            text=text,
            success=True,
            error_message=None,
            file_type=file_type,
            word_count=len(words),
            char_count=len(text),
        )

    except Exception as exc:
        return ResumeParseResult(
            text="",
            success=False,
            error_message=str(exc),
            file_type=ext.replace(".", "") or "unknown",
            word_count=0,
            char_count=0,
        )
