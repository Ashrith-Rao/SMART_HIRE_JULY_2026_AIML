"""Comprehensive unit and integration test suite for SmartHire ML & NLP pipelines."""

import io
from pathlib import Path
import docx
import numpy as np
import pandas as pd
import pytest

from src.config import CLASSIFIER_PATH, JOB_METADATA_PATH
from src.features.match_features import compute_composite_fit_score, compute_skill_gap, parse_experience_range
from src.features.text_features import clean_text, extract_skills, normalize_skill
from src.models.classifier import ResumeClassifier
from src.models.recommender import JobRecommender
from src.parsing.resume_parser import parse_resume


# =========================================================================
# 1. TEXT CLEANING & TOKEN PRESERVATION TESTS
# =========================================================================

def test_clean_text_basic():
    raw = "Hello World! Check out https://github.com/test and email us at test@example.com."
    cleaned = clean_text(raw)
    assert "https" not in cleaned
    assert "github.com" not in cleaned
    assert "test@example.com" not in cleaned
    assert "hello world" in cleaned


def test_clean_text_preserves_technical_tokens():
    raw = "Experienced in C++, C#, .NET, Node.js, and React.js with CI/CD pipelines."
    cleaned = clean_text(raw, preserve_tech_tokens=True)
    assert "c++" in cleaned
    assert "c#" in cleaned
    assert ".net" in cleaned
    assert "node.js" in cleaned
    assert "react.js" in cleaned
    assert "ci/cd" in cleaned


def test_clean_text_empty_and_short():
    assert clean_text("") == ""
    assert clean_text("   ") == ""
    assert clean_text(None) == ""


# =========================================================================
# 2. SKILL EXTRACTION & NORMALIZATION TESTS
# =========================================================================

def test_normalize_skill_cases():
    assert normalize_skill("Python") == "python"
    assert normalize_skill("PYTHON") == "python"
    assert normalize_skill("python") == "python"
    assert normalize_skill("Scikit Learn") == "scikit-learn"
    assert normalize_skill("sklearn") == "scikit-learn"
    assert normalize_skill("ML") == "machine learning"
    assert normalize_skill("JS") == "javascript"
    assert normalize_skill("AWS") == "aws"
    assert normalize_skill("Amazon Web Services") == "aws"


def test_extract_skills_technical_tokens():
    text = "Full stack developer with 4 years in C++, C#, .NET, Node.js, React.js, and Docker."
    skills = extract_skills(text)
    assert "c++" in skills
    assert "c#" in skills
    assert ".net" in skills
    assert "node.js" in skills
    assert "react.js" in skills
    assert "docker" in skills


def test_extract_skills_boundary_safety():
    # Ensure single letters like 'r' or short words don't trigger false positives
    text = "Our cat ran on the road to the store."
    skills = extract_skills(text)
    assert "r" not in skills
    assert "c" not in skills


def test_extract_skills_data_science():
    text = "Experienced Data Scientist skilled in Python, SQL, Pandas, Scikit-learn, Machine Learning, and Tableau."
    skills = extract_skills(text)
    for expected in ["python", "sql", "pandas", "scikit-learn", "machine learning", "tableau"]:
        assert expected in skills


# =========================================================================
# 3. SKILL GAP COMPUTATION TESTS
# =========================================================================

def test_compute_skill_gap_exact():
    candidate_skills = ["Python", "SQL", "Pandas", "Machine Learning"]
    job_skills = ["Python", "SQL", "Pandas", "Tableau", "Power BI"]

    report = compute_skill_gap(candidate_skills, job_skills)
    assert report["matched_skills"] == ["pandas", "python", "sql"]
    assert report["missing_skills"] == ["power bi", "tableau"]
    # 3 matched out of 5 required = 60.0%
    assert report["skill_coverage_pct"] == 60.0


def test_compute_skill_gap_all_matched():
    candidate_skills = ["python", "sql", "aws", "docker"]
    job_skills = ["python", "sql"]
    report = compute_skill_gap(candidate_skills, job_skills)
    assert report["skill_coverage_pct"] == 100.0
    assert len(report["missing_skills"]) == 0


def test_compute_skill_gap_empty_inputs():
    report_empty_job = compute_skill_gap(["python"], [])
    assert report_empty_job["skill_coverage_pct"] == 100.0

    report_empty_cand = compute_skill_gap([], ["python", "sql"])
    assert report_empty_cand["skill_coverage_pct"] == 0.0
    assert report_empty_cand["missing_skills"] == ["python", "sql"]


# =========================================================================
# 4. EXPERIENCE PARSER & FIT SCORE TESTS
# =========================================================================

def test_parse_experience_range():
    min_e, max_e = parse_experience_range("2 - 5 yrs")
    assert min_e == 2.0
    assert max_e == 5.0

    min_e, max_e = parse_experience_range("5 yrs")
    assert min_e == 5.0

    min_e, max_e = parse_experience_range("")
    assert min_e == 0.0
    assert max_e == 30.0


def test_compute_composite_fit_score():
    score_dict = compute_composite_fit_score(
        text_similarity=0.80,
        skill_coverage_pct=75.0,
        candidate_exp_years=4.0,
        job_exp_str="2-5 yrs",
        candidate_category="Data Science",
        job_category="Data Scientist",
    )
    assert 0.0 <= score_dict["composite_score"] <= 100.0
    assert "fit_tier" in score_dict
    assert "components" in score_dict
    assert score_dict["components"]["text_similarity_pct"] == 80.0
    assert score_dict["components"]["skill_coverage_pct"] == 75.0


# =========================================================================
# 5. RESUME PARSER TESTS
# =========================================================================

def test_parse_resume_txt(tmp_path):
    txt_file = tmp_path / "sample_resume.txt"
    sample_content = (
        "John Doe\n"
        "Senior Software Engineer\n"
        "Skills: Python, Django, PostgreSQL, Docker, AWS, Git.\n"
        "Experience: 5 years designing scalable web microservices."
    )
    txt_file.write_text(sample_content, encoding="utf-8")

    result = parse_resume(txt_file)
    assert result.success is True
    assert result.file_type == "txt"
    assert "John Doe" in result.text
    assert result.word_count > 10


def test_parse_resume_docx(tmp_path):
    doc_file = tmp_path / "sample_resume.docx"
    doc = docx.Document()
    doc.add_heading("Jane Smith — Data Scientist", level=1)
    doc.add_paragraph("Core competencies: Machine Learning, Python, PyTorch, SQL, Statistics.")
    doc.save(str(doc_file))

    result = parse_resume(doc_file)
    assert result.success is True
    assert result.file_type == "docx"
    assert "Jane Smith" in result.text
    assert "PyTorch" in result.text


def test_parse_resume_pdf():
    pdf_file = Path("data/sample_resumes/sample_data_scientist.pdf")
    if not pdf_file.exists():
        pytest.skip("Sample PDF not found.")
    result = parse_resume(pdf_file)
    assert result.success is True
    assert result.file_type == "pdf"
    assert "Data Scientist" in result.text
    assert "Python" in result.text
    assert result.word_count > 10


def test_parse_resume_empty_file(tmp_path):
    empty_file = tmp_path / "empty.txt"
    empty_file.write_text("", encoding="utf-8")
    result = parse_resume(empty_file)
    assert result.success is False
    assert "empty" in result.error_message.lower()


def test_parse_resume_non_existent_file():
    result = parse_resume(Path("non_existent_folder/fake_resume.pdf"))
    assert result.success is False
    assert "not exist" in result.error_message.lower()


def test_parse_resume_unsupported_format(tmp_path):
    bad_file = tmp_path / "resume.zip"
    bad_file.write_bytes(b"PK000fakezipcontent")
    result = parse_resume(bad_file)
    assert result.success is False
    assert "unsupported" in result.error_message.lower()


# =========================================================================
# 6. MODEL INFERENCE & ARTIFACT TESTS
# =========================================================================

def test_classifier_loading_and_inference():
    if not CLASSIFIER_PATH.exists():
        pytest.skip("Classifier artifact not yet trained.")

    clf = ResumeClassifier.load()
    assert clf.model is not None
    assert clf.vectorizer is not None

    ds_resume = (
        "Experienced Data Scientist with 4 years in Python, machine learning, deep learning, "
        "Scikit-learn, TensorFlow, SQL, pandas, numpy, and predictive modeling."
    )
    pred = clf.predict(ds_resume)
    assert "predicted_category" in pred
    assert pred["confidence"] > 0.3
    assert len(pred["top_categories"]) >= 1
    assert "data science" in pred["predicted_category"].lower() or "python" in pred["predicted_category"].lower()


def test_recommender_loading_and_ranking():
    if not JOB_METADATA_PATH.exists():
        pytest.skip("Recommender artifacts not yet built.")

    rec = JobRecommender.load()
    assert rec.vectorizer is not None
    assert rec.job_matrix is not None
    assert len(rec.jobs_df) > 0

    resume_text = (
        "Full stack Python developer with experience in Django, React.js, PostgreSQL, Docker, and REST APIs."
    )
    recommendations = rec.recommend(resume_text, top_n=5)
    assert len(recommendations) == 5
    # Check descending order of similarity scores
    scores = [r["similarity_score"] for r in recommendations]
    assert scores == sorted(scores, reverse=True)

    # Check deduplication
    pairs = [(r["title"].lower(), r["company"].lower()) for r in recommendations]
    assert len(pairs) == len(set(pairs))

    # Check structure of each recommendation
    first = recommendations[0]
    assert "rank" in first
    assert "title" in first
    assert "company" in first
    assert "matched_skills" in first
    assert "missing_skills" in first
    assert "composite_fit" in first
    assert "explanation" in first
